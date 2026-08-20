#!/usr/bin/env python3
"""QA 知识库 Excel 协作、校验与 SQLite 安全导入工具。

工作流：
1. 伙伴只编辑 ``qa/qa_content.xlsx``；
2. 运行 ``sync-csv`` 生成可供 Git/PR 审查的 CSV 差异；
3. 先用 ``import-db`` 默认的 dry-run 查看变更；
4. 只有显式加 ``--apply`` 才会写入，写入前自动备份 SQLite。

为了避免误操作正式库，所有数据库命令都必须显式传 ``--db``。
"""
from __future__ import annotations

import argparse
import csv
import json
import sqlite3
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import openpyxl


SHEET_NAME = "QA内容"
MAX_ROWS = 5000
MAX_QUESTION = 1000
MAX_ANSWER = 10000
ALLOWED_ACTIONS = {"KEEP", "ADD", "DISABLE"}

COLUMNS = [
    ("action", "操作"),
    ("id", "ID"),
    ("question", "问题"),
    ("answer", "标准答案"),
    ("answer_variants", "备用答案(JSON)"),
    ("category", "分类"),
    ("tags", "标签"),
    ("domains", "业务域"),
    ("status", "状态"),
    ("enabled", "启用"),
    ("version", "版本"),
    ("change_note", "修改说明"),
]
KEYS = [key for key, _ in COLUMNS]
HEADERS = [header for _, header in COLUMNS]


class WorkflowError(RuntimeError):
    pass


@dataclass
class QaRow:
    row_no: int
    action: str
    id: int | None
    question: str
    answer: str
    answer_variants: str
    category: str
    tags: str
    domains: str
    status: str
    enabled: str
    version: int | None
    change_note: str

    def as_dict(self) -> dict[str, Any]:
        return {key: getattr(self, key) for key in KEYS}


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _integer(value: Any, label: str, row_no: int, required: bool = False) -> int | None:
    text = _text(value)
    if not text:
        if required:
            raise WorkflowError(f"第 {row_no} 行：{label} 不能为空")
        return None
    try:
        number = int(float(text))
    except (TypeError, ValueError) as exc:
        raise WorkflowError(f"第 {row_no} 行：{label} 必须是整数，当前为 {text!r}") from exc
    if number < 1:
        raise WorkflowError(f"第 {row_no} 行：{label} 必须大于 0")
    return number


def _normalize_variants(value: Any, row_no: int) -> str:
    text = _text(value)
    if not text:
        return ""
    try:
        if text.startswith("["):
            parsed = json.loads(text)
        else:
            parsed = [line.strip() for line in text.splitlines() if line.strip()]
    except json.JSONDecodeError as exc:
        raise WorkflowError(f"第 {row_no} 行：备用答案不是有效 JSON 数组") from exc
    if not isinstance(parsed, list) or not all(isinstance(item, str) for item in parsed):
        raise WorkflowError(f"第 {row_no} 行：备用答案必须是字符串数组")
    unique: list[str] = []
    for item in parsed:
        item = item.strip()
        if item and item not in unique:
            if len(item) > MAX_ANSWER:
                raise WorkflowError(f"第 {row_no} 行：备用答案超过 {MAX_ANSWER} 字")
            unique.append(item)
    return json.dumps(unique, ensure_ascii=False) if unique else ""


def load_xlsx(path: Path) -> list[QaRow]:
    if not path.exists():
        raise WorkflowError(f"Excel 不存在：{path}")
    # artifact-tool 生成的工作表可以不带 openpyxl 只读模式依赖的
    # dimension 元数据；普通模式会从实际单元格正确计算行数。
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
    if SHEET_NAME not in workbook.sheetnames:
        raise WorkflowError(f"Excel 缺少工作表：{SHEET_NAME}")
    sheet = workbook[SHEET_NAME]
    header_row_no = None
    header_values: list[str] = []
    for row in sheet.iter_rows(min_row=1, max_row=min(12, sheet.max_row or 12)):
        values = [_text(cell.value) for cell in row]
        if values[: len(HEADERS)] == HEADERS:
            header_row_no = row[0].row
            header_values = values
            break
    if header_row_no is None:
        raise WorkflowError(f"Excel 表头不匹配，应为：{', '.join(HEADERS)}")
    if header_values[: len(HEADERS)] != HEADERS:
        raise WorkflowError("Excel 表头顺序被修改，请恢复模板")

    rows: list[QaRow] = []
    for cells in sheet.iter_rows(min_row=header_row_no + 1, max_col=len(HEADERS)):
        row_no = cells[0].row
        if any(cell.data_type == "f" for cell in cells):
            raise WorkflowError(f"第 {row_no} 行：不允许使用 Excel 公式")
        values = [cell.value for cell in cells]
        action = (_text(values[0]) or "KEEP").upper()
        id_value = _integer(values[1], "ID", row_no)
        question = _text(values[2])
        answer = _text(values[3])
        # 预留的空白新增行不参与校验/导入。
        if id_value is None and not question and not answer:
            continue
        rows.append(QaRow(
            row_no=row_no,
            action=action,
            id=id_value,
            question=question,
            answer=answer,
            answer_variants=_normalize_variants(values[4], row_no),
            category=_text(values[5]),
            tags=_text(values[6]),
            domains=_text(values[7]),
            status=_text(values[8]) or "approved",
            enabled=_text(values[9]) or "1",
            version=_integer(values[10], "版本", row_no),
            change_note=_text(values[11]),
        ))
    validate_rows(rows)
    return rows


def validate_rows(rows: list[QaRow]) -> None:
    errors: list[str] = []
    if not rows:
        errors.append("Excel 中没有 QA 数据")
    if len(rows) > MAX_ROWS:
        errors.append(f"QA 行数 {len(rows)} 超过上限 {MAX_ROWS}")

    seen_ids: dict[int, int] = {}
    seen_questions: dict[str, int] = {}
    for row in rows:
        if row.action not in ALLOWED_ACTIONS:
            errors.append(f"第 {row.row_no} 行：操作只能是 KEEP / ADD / DISABLE")
        if row.action == "ADD" and row.id is not None:
            errors.append(f"第 {row.row_no} 行：ADD 行的 ID 必须留空")
        if row.action in {"KEEP", "DISABLE"} and row.id is None:
            errors.append(f"第 {row.row_no} 行：{row.action} 行必须保留 ID")
        if row.action in {"KEEP", "DISABLE"} and row.version is None:
            errors.append(f"第 {row.row_no} 行：{row.action} 行必须保留版本")
        if not row.question:
            errors.append(f"第 {row.row_no} 行：问题不能为空")
        elif len(row.question) > MAX_QUESTION:
            errors.append(f"第 {row.row_no} 行：问题超过 {MAX_QUESTION} 字")
        if not row.answer:
            errors.append(f"第 {row.row_no} 行：标准答案不能为空")
        elif len(row.answer) > MAX_ANSWER:
            errors.append(f"第 {row.row_no} 行：标准答案超过 {MAX_ANSWER} 字")
        if row.enabled not in {"0", "1"}:
            errors.append(f"第 {row.row_no} 行：启用只能是 0 或 1")
        if row.id is not None:
            if row.id in seen_ids:
                errors.append(f"第 {row.row_no} 行：ID {row.id} 与第 {seen_ids[row.id]} 行重复")
            seen_ids[row.id] = row.row_no
        if row.action != "DISABLE" and row.question:
            key = row.question.casefold()
            if key in seen_questions:
                errors.append(f"第 {row.row_no} 行：问题与第 {seen_questions[key]} 行完全重复")
            seen_questions[key] = row.row_no

    if errors:
        preview = "\n".join(f"- {item}" for item in errors[:30])
        more = f"\n- 还有 {len(errors) - 30} 个错误" if len(errors) > 30 else ""
        raise WorkflowError(f"QA 内容校验失败：\n{preview}{more}")


def _csv_safe(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    return "'" + value if value.startswith(("=", "+", "-", "@", "\t", "\r")) else value


def write_csv(rows: Iterable[QaRow | dict[str, Any]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=HEADERS,
            quoting=csv.QUOTE_ALL,
            lineterminator="\n",
        )
        writer.writeheader()
        for item in rows:
            raw = item.as_dict() if isinstance(item, QaRow) else item
            writer.writerow({header: _csv_safe(raw.get(key, "")) for key, header in COLUMNS})


def _table_columns(connection: sqlite3.Connection) -> set[str]:
    rows = connection.execute("PRAGMA table_info(qa_item)").fetchall()
    if not rows:
        raise WorkflowError("数据库中不存在 qa_item 表")
    return {row[1] for row in rows}


def export_source(db_path: Path, json_path: Path, csv_path: Path) -> None:
    if not db_path.exists():
        raise WorkflowError(f"数据库不存在：{db_path}")
    connection = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    columns = _table_columns(connection)
    optional = [name for name in ("answer_variants", "domains") if name in columns]
    select_columns = [
        "id", "question", "answer", *optional, "category", "tags",
        "status", "enabled", "version", "updated_at",
    ]
    query = f"SELECT {', '.join(select_columns)} FROM qa_item WHERE deleted='0' ORDER BY id"
    exported: list[dict[str, Any]] = []
    for record in connection.execute(query):
        item = dict(record)
        variants = _normalize_variants(item.get("answer_variants", ""), int(item["id"]))
        exported.append({
            "action": "KEEP",
            "id": item["id"],
            "question": item.get("question") or "",
            "answer": item.get("answer") or "",
            "answer_variants": variants,
            "category": item.get("category") or "",
            "tags": item.get("tags") or "",
            "domains": item.get("domains") or "",
            "status": item.get("status") or "approved",
            "enabled": item.get("enabled") or "1",
            "version": item.get("version") or 1,
            "change_note": "",
        })
    connection.close()
    json_path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "sheet_name": SHEET_NAME,
        "headers": HEADERS,
        "columns": KEYS,
        "exported_at": datetime.now(timezone.utc).isoformat(),
        "source_db": db_path.name,
        "row_count": len(exported),
        "rows": exported,
    }
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    write_csv(exported, csv_path)
    print(f"已导出 {len(exported)} 条 QA -> {csv_path}")


def _normalized(value: Any) -> str:
    return "" if value is None else str(value).strip()


def plan_changes(
    connection: sqlite3.Connection,
    rows: list[QaRow],
    approval_mode: str,
    actor: str,
) -> tuple[list[tuple[str, tuple[Any, ...], str]], dict[str, int]]:
    columns = _table_columns(connection)
    connection.row_factory = sqlite3.Row
    current_by_id = {
        int(row["id"]): dict(row)
        for row in connection.execute("SELECT * FROM qa_item")
    }
    current_by_question = {
        _normalized(row["question"]).casefold(): dict(row)
        for row in current_by_id.values()
        if _normalized(row.get("deleted", "0")) == "0"
    }
    now = datetime.now(timezone.utc).replace(tzinfo=None).isoformat(sep=" ")
    operations: list[tuple[str, tuple[Any, ...], str]] = []
    counts = {"add": 0, "update": 0, "disable": 0, "unchanged": 0, "duplicate": 0}
    conflicts: list[str] = []
    content_fields = [name for name in ("question", "answer", "answer_variants", "category", "tags", "domains") if name in columns]

    for row in rows:
        if row.action == "ADD":
            existing = current_by_question.get(row.question.casefold())
            if existing:
                counts["duplicate"] += 1
                continue
            insert_fields = ["question", "answer", "category", "tags"]
            if "answer_variants" in columns:
                insert_fields.append("answer_variants")
            if "domains" in columns:
                insert_fields.append("domains")
            insert_fields += ["status", "enabled", "deleted", "version", "created_by", "created_at", "updated_by", "updated_at"]
            values: dict[str, Any] = {
                "question": row.question,
                "answer": row.answer,
                "answer_variants": row.answer_variants or None,
                "category": row.category or None,
                "tags": row.tags or None,
                "domains": row.domains or None,
                "status": approval_mode,
                "enabled": "1",
                "deleted": "0",
                "version": 1,
                "created_by": actor,
                "created_at": now,
                "updated_by": actor,
                "updated_at": now,
            }
            placeholders = ", ".join("?" for _ in insert_fields)
            sql = f"INSERT INTO qa_item ({', '.join(insert_fields)}) VALUES ({placeholders})"
            operations.append((sql, tuple(values[field] for field in insert_fields), f"ADD: {row.question[:50]}"))
            counts["add"] += 1
            continue

        current = current_by_id.get(int(row.id or 0))
        if not current:
            conflicts.append(f"第 {row.row_no} 行：目标库不存在 ID {row.id}")
            continue
        if row.action == "DISABLE":
            current_version = int(current.get("version") or 1)
            if row.version != current_version:
                conflicts.append(
                    f"第 {row.row_no} 行：ID {row.id} 版本冲突（Excel={row.version}, 数据库={current_version}）"
                )
                continue
            if _normalized(current.get("enabled", "1")) == "0":
                counts["unchanged"] += 1
                continue
            next_version = current_version + 1
            sql = "UPDATE qa_item SET enabled='0', status='disabled', version=?, updated_by=?, updated_at=? WHERE id=?"
            operations.append((sql, (next_version, actor, now, row.id), f"DISABLE ID {row.id}: {row.question[:40]}"))
            counts["disable"] += 1
            continue

        row_values = row.as_dict()
        changed_fields = [
            field for field in content_fields
            if _normalized(current.get(field)) != _normalized(row_values.get(field))
        ]
        if not changed_fields:
            counts["unchanged"] += 1
            continue
        current_version = int(current.get("version") or 1)
        if row.version != current_version:
            conflicts.append(
                f"第 {row.row_no} 行：ID {row.id} 版本冲突（Excel={row.version}, 数据库={current_version}）"
            )
            continue
        set_fields = changed_fields + ["status", "enabled", "version", "updated_by", "updated_at"]
        values = [row_values[field] or None for field in changed_fields]
        values += [approval_mode, "1", current_version + 1, actor, now, row.id]
        assignments = ", ".join(f"{field}=?" for field in set_fields)
        sql = f"UPDATE qa_item SET {assignments} WHERE id=?"
        operations.append((sql, tuple(values), f"UPDATE ID {row.id}: {', '.join(changed_fields)}"))
        counts["update"] += 1

    if conflicts:
        raise WorkflowError("导入前发现冲突：\n" + "\n".join(f"- {item}" for item in conflicts[:30]))
    return operations, counts


def import_db(
    db_path: Path,
    xlsx_path: Path,
    apply: bool,
    approval_mode: str,
    actor: str,
) -> None:
    if not db_path.exists():
        raise WorkflowError(f"数据库不存在：{db_path}")
    rows = load_xlsx(xlsx_path)
    connection = sqlite3.connect(db_path)
    operations, counts = plan_changes(connection, rows, approval_mode, actor)
    print(
        f"变更计划：新增 {counts['add']}，更新 {counts['update']}，"
        f"停用 {counts['disable']}，无变化 {counts['unchanged']}，重复跳过 {counts['duplicate']}"
    )
    for _, _, summary in operations[:40]:
        print(f"  - {summary}")
    if len(operations) > 40:
        print(f"  - ... 还有 {len(operations) - 40} 项")
    if not apply:
        print("DRY-RUN：未写入数据库。确认后加 --apply。")
        connection.close()
        return

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup_path = db_path.with_name(f"{db_path.name}.bak.{stamp}")
    backup = sqlite3.connect(backup_path)
    connection.backup(backup)
    backup.close()
    try:
        connection.execute("BEGIN IMMEDIATE")
        for sql, params, _ in operations:
            connection.execute(sql, params)
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
    print(f"已写入 {len(operations)} 项变更；备份：{backup_path}")


QA_SCHEMA = """
CREATE TABLE qa_item (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  question VARCHAR(1000) NOT NULL,
  answer TEXT NOT NULL,
  answer_variants TEXT,
  category VARCHAR(64),
  tags VARCHAR(500),
  domains VARCHAR(255),
  status VARCHAR(16) DEFAULT 'approved',
  version INTEGER DEFAULT 1,
  enabled VARCHAR(1) DEFAULT '1',
  deleted VARCHAR(1) DEFAULT '0',
  created_by VARCHAR(64),
  created_at DATETIME,
  updated_by VARCHAR(64),
  updated_at DATETIME,
  approved_by VARCHAR(64),
  approved_at DATETIME
);
CREATE INDEX ix_qa_alive ON qa_item(deleted, status, enabled);
CREATE INDEX ix_qa_updated_at ON qa_item(updated_at);
"""


def init_test_db(db_path: Path, xlsx_path: Path) -> None:
    if db_path.exists():
        raise WorkflowError(f"为避免覆盖，目标已存在：{db_path}")
    rows = load_xlsx(xlsx_path)
    db_path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(db_path)
    try:
        connection.executescript(QA_SCHEMA)
        now = datetime.now(timezone.utc).replace(tzinfo=None).isoformat(sep=" ")
        inserted = 0
        for row in rows:
            if row.action == "DISABLE":
                continue
            fields = [
                "question", "answer", "answer_variants", "category", "tags", "domains",
                "status", "version", "enabled", "deleted", "created_by", "created_at",
                "updated_by", "updated_at", "approved_by", "approved_at",
            ]
            values: list[Any] = [
                row.question, row.answer, row.answer_variants or None, row.category or None,
                row.tags or None, row.domains or None, "approved", row.version or 1,
                "1", "0", "qa-content-workbook", now, "qa-content-workbook", now,
                "qa-content-workbook", now,
            ]
            if row.id is not None:
                fields.insert(0, "id")
                values.insert(0, row.id)
            placeholders = ", ".join("?" for _ in fields)
            connection.execute(
                f"INSERT INTO qa_item ({', '.join(fields)}) VALUES ({placeholders})",
                values,
            )
            inserted += 1
        connection.commit()
    except Exception:
        connection.rollback()
        connection.close()
        if db_path.exists():
            db_path.unlink()
        raise
    connection.close()
    print(f"已创建测试 QA 库：{db_path}（{inserted} 条，全部 approved）")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)

    export = sub.add_parser("export-source", help="从 SQLite 导出 builder JSON 和审查 CSV")
    export.add_argument("--db", type=Path, required=True)
    export.add_argument("--json", type=Path, required=True)
    export.add_argument("--csv", type=Path, required=True)

    check = sub.add_parser("check", help="校验 QA Excel")
    check.add_argument("--xlsx", type=Path, default=Path("qa/qa_content.xlsx"))

    sync = sub.add_parser("sync-csv", help="校验 Excel 并生成 Git 可读 CSV")
    sync.add_argument("--xlsx", type=Path, default=Path("qa/qa_content.xlsx"))
    sync.add_argument("--csv", type=Path, default=Path("qa/qa_content.csv"))

    import_parser = sub.add_parser("import-db", help="预览或导入到显式指定的 SQLite")
    import_parser.add_argument("--db", type=Path, required=True)
    import_parser.add_argument("--xlsx", type=Path, default=Path("qa/qa_content.xlsx"))
    import_parser.add_argument("--apply", action="store_true")
    import_parser.add_argument("--approval-mode", choices=("pending", "approved"), default="pending")
    import_parser.add_argument("--actor", default="qa-content-import")

    init = sub.add_parser("init-test-db", help="用 Excel 新建一个独立 QA 测试库")
    init.add_argument("--db", type=Path, required=True)
    init.add_argument("--xlsx", type=Path, default=Path("qa/qa_content.xlsx"))
    return parser


def main() -> int:
    args = build_parser().parse_args()
    try:
        if args.command == "export-source":
            export_source(args.db, args.json, args.csv)
        elif args.command == "check":
            rows = load_xlsx(args.xlsx)
            print(f"校验通过：{len(rows)} 条 QA")
        elif args.command == "sync-csv":
            rows = load_xlsx(args.xlsx)
            write_csv(rows, args.csv)
            print(f"校验通过：{len(rows)} 条 QA -> {args.csv}")
        elif args.command == "import-db":
            import_db(args.db, args.xlsx, args.apply, args.approval_mode, args.actor)
        elif args.command == "init-test-db":
            init_test_db(args.db, args.xlsx)
        return 0
    except WorkflowError as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
