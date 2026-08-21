#!/usr/bin/env python3
"""把正式 SQLite 安全复制成完整测试快照，并移除所有正式出向凭证。"""
from __future__ import annotations

import argparse
import shutil
import sqlite3
import tempfile
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill


HEADERS = [
    "操作", "ID", "问题", "标准答案", "备用答案(JSON)", "分类",
    "标签", "业务域", "状态", "启用", "版本", "修改说明",
]

SECRET_SETTING_KEYS = {
    "admin_url",
    "alert_secret",
    "alert_webhook",
    "ali_bailian_dingtalk_webhook",
    "ali_bailian_dingtalk_reply_mobile",
    "ali_bailian_dingtalk_reply_staff_id",
    "dingtalk_client_id",
    "dingtalk_client_secret",
    "dingtalk_robot_code",
    "feishu_app_id",
    "feishu_app_secret",
    "feishu_bitable_app_token",
    "feishu_bitable_table_id",
    "feishu_oauth_redirect_base",
    "llm_api_key",
    "moderation_api_key",
    "public_base_url",
    "vision_api_key",
    "volcano_bitable_app_token",
    "volcano_bitable_table_id",
}

DISABLED_SETTING_KEYS = {
    "daily_brief_enabled",
    "failure_replay_enabled",
    "feishu_enabled",
    "llm_enabled",
    "moderation_query_enabled",
    "vision_enabled",
    "volcano_ticket_poll_enabled",
    "wecom_bridge_read_enabled",
    "wecom_bridge_send_enabled",
}


def _tables(connection: sqlite3.Connection) -> set[str]:
    return {
        str(row[0])
        for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")
    }


def sanitize_database(path: Path) -> dict[str, int]:
    """保留业务内容，移除会触发正式外部系统的凭证和启用状态。"""
    connection = sqlite3.connect(path)
    tables = _tables(connection)
    counts: dict[str, int] = {}
    try:
        if "sys_setting" in tables:
            placeholders = ",".join("?" for _ in SECRET_SETTING_KEYS)
            cursor = connection.execute(
                f"UPDATE sys_setting SET v='' WHERE k IN ({placeholders})",
                tuple(sorted(SECRET_SETTING_KEYS)),
            )
            counts["cleared_settings"] = max(cursor.rowcount, 0)
            placeholders = ",".join("?" for _ in DISABLED_SETTING_KEYS)
            cursor = connection.execute(
                f"UPDATE sys_setting SET v='False' WHERE k IN ({placeholders})",
                tuple(sorted(DISABLED_SETTING_KEYS)),
            )
            counts["disabled_settings"] = max(cursor.rowcount, 0)

        if "dingtalk_group" in tables:
            cursor = connection.execute(
                """
                UPDATE dingtalk_group
                   SET active='0', webhook_url=NULL, webhook_secret=NULL, robot_code=NULL,
                       note=CASE
                         WHEN COALESCE(note, '')='' THEN '正式环境快照（测试环境已禁用）'
                         ELSE note || ' | 正式环境快照（测试环境已禁用）'
                       END
                """
            )
            counts["disabled_groups"] = max(cursor.rowcount, 0)

        if "broadcast_schedule" in tables:
            cursor = connection.execute("UPDATE broadcast_schedule SET enabled='0'")
            counts["disabled_schedules"] = max(cursor.rowcount, 0)

        if "uploaded_file" in tables:
            connection.execute(
                """
                UPDATE uploaded_file
                   SET public_url=CASE
                         WHEN COALESCE(filename, '')='' THEN NULL
                         ELSE '/files/' || filename
                       END,
                       dingtalk_media_id=NULL,
                       dingtalk_media_uploaded_at=NULL
                """
            )

        if "sys_user" in tables:
            # 服务重启时只会重新启用测试 owner/partner 两个强密码账号。
            cursor = connection.execute("UPDATE sys_user SET enabled='0'")
            counts["disabled_users"] = max(cursor.rowcount, 0)

        connection.commit()
    finally:
        connection.close()
    return counts


def write_baseline_xlsx(db_path: Path, output: Path) -> int:
    """从本次正式快照生成 QA 差异基线，避免把已有正式 QA 误判为新增。"""
    connection = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    try:
        rows = connection.execute(
            """
            SELECT id, question, answer, answer_variants, category, tags, domains,
                   status, enabled, version
              FROM qa_item
             WHERE COALESCE(deleted, '0')='0' AND COALESCE(enabled, '1')='1'
             ORDER BY id
            """
        ).fetchall()
    finally:
        connection.close()

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "QA内容"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
    for row in rows:
        sheet.append([
            "KEEP", row["id"], row["question"] or "", row["answer"] or "",
            row["answer_variants"] or "", row["category"] or "", row["tags"] or "",
            row["domains"] or "", row["status"] or "approved", row["enabled"] or "1",
            row["version"] or 1, "",
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return len(rows)


def snapshot_database(source: Path, destination: Path) -> Path | None:
    source = source.expanduser().resolve()
    destination = destination.expanduser().resolve()
    if source == destination:
        raise RuntimeError("正式库和测试库不能是同一个文件")
    if not any(word in destination.name.lower() for word in ("workspace", "staging", "test")):
        raise RuntimeError("为防止覆盖正式库，目标文件名必须包含 workspace、staging 或 test")
    if not source.exists():
        raise RuntimeError(f"正式库不存在：{source}")

    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix=f".{destination.stem}-next-", suffix=".db", dir=destination.parent, delete=False,
    ) as handle:
        pending = Path(handle.name)
    try:
        source_connection = sqlite3.connect(f"file:{source}?mode=ro", uri=True)
        target_connection = sqlite3.connect(pending)
        try:
            source_connection.backup(target_connection)
        finally:
            target_connection.close()
            source_connection.close()
        sanitize_database(pending)

        backup = None
        if destination.exists():
            stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
            backup = destination.with_name(f"{destination.name}.before-refresh-{stamp}")
            destination.replace(backup)
        try:
            pending.replace(destination)
        except Exception:
            if backup and backup.exists() and not destination.exists():
                backup.replace(destination)
            raise
        return backup
    finally:
        if pending.exists():
            pending.unlink()


def replace_uploads(source: Path, destination: Path) -> Path | None:
    if not source.exists():
        return None
    temp_dir = Path(tempfile.mkdtemp(prefix="qa-staging-uploads-", dir=destination.parent))
    shutil.rmtree(temp_dir)
    shutil.copytree(source, temp_dir, symlinks=True)
    backup = None
    if destination.exists():
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        backup = destination.with_name(f"{destination.name}.before-refresh-{stamp}")
        destination.replace(backup)
    temp_dir.replace(destination)
    return backup


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-db", type=Path, required=True)
    parser.add_argument("--target-db", type=Path, required=True)
    parser.add_argument("--baseline-xlsx", type=Path, required=True)
    parser.add_argument("--source-uploads", type=Path)
    parser.add_argument("--target-uploads", type=Path)
    args = parser.parse_args()

    backup = snapshot_database(args.source_db, args.target_db)
    baseline_count = write_baseline_xlsx(args.target_db.resolve(), args.baseline_xlsx.resolve())
    upload_backup = None
    if args.source_uploads and args.target_uploads:
        args.target_uploads.parent.mkdir(parents=True, exist_ok=True)
        upload_backup = replace_uploads(args.source_uploads.resolve(), args.target_uploads.resolve())

    connection = sqlite3.connect(args.target_db)
    try:
        tables = _tables(connection)
        summary = {
            table: connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
            for table in ("qa_item", "conversation", "dingtalk_group", "broadcast", "uploaded_file")
            if table in tables
        }
    finally:
        connection.close()
    print(f"测试快照已生成：{args.target_db.resolve()}")
    print(f"QA 差异基线：{baseline_count} 条 -> {args.baseline_xlsx.resolve()}")
    print(f"内容统计：{summary}")
    if backup:
        print(f"旧测试库备份：{backup}")
    if upload_backup:
        print(f"旧测试附件备份：{upload_backup}")


if __name__ == "__main__":
    main()
