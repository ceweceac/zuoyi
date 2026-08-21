import io
import sqlite3
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

import openpyxl

from backend.app.services.qa_workspace import (
    HEADERS,
    WorkspaceChange,
    build_change_package,
    collect_changes,
    load_baseline,
)
from scripts.refresh_staging_from_production import sanitize_database, write_baseline_xlsx


def item(**values):
    defaults = {
        "id": 1,
        "question": "问题",
        "answer": "答案",
        "answer_variants": "",
        "category": "分类",
        "tags": "",
        "domains": "",
        "status": "approved",
        "enabled": "1",
        "deleted": "0",
        "version": 1,
        "created_by": "baseline",
        "updated_by": "baseline",
    }
    defaults.update(values)
    return SimpleNamespace(**defaults)


class QaWorkspaceTests(unittest.TestCase):
    def make_baseline(self, path: Path):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "QA内容"
        sheet.append(HEADERS)
        sheet.append(["KEEP", 1, "问题一", "答案一", "", "A", "", "", "approved", "1", 3, ""])
        sheet.append(["KEEP", 2, "问题二", "答案二", "", "B", "", "", "approved", "1", 7, ""])
        sheet.append(["KEEP", 3, "问题三", "答案三", "", "C", "", "", "approved", "1", 2, ""])
        workbook.save(path)

    def test_load_baseline_and_collect_reviewed_changes(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "baseline.xlsx"
            self.make_baseline(path)
            baseline = load_baseline(path)

        changes, stats = collect_changes([
            item(id=1, question="问题一", answer="更新答案", category="A", version=4, updated_by="qa_partner"),
            item(id=2, question="问题二", answer="答案二", category="B", enabled="0", version=8, updated_by="qa_owner"),
            item(id=3, question="问题三", answer="待审核答案", category="C", status="pending", version=3, updated_by="qa_partner"),
            item(id=4, question="新增问题", answer="新增答案", category="D", status="approved", updated_by="qa_partner"),
        ], baseline)

        self.assertEqual(stats, {"add": 1, "update": 1, "disable": 1, "pending": 1, "ready": 3})
        self.assertEqual([change.action for change in changes], ["KEEP", "DISABLE", "ADD"])
        # 导出必须使用正式基线版本，而不是测试库已自增的版本。
        self.assertEqual(changes[0].version, 3)
        self.assertEqual(changes[1].version, 7)
        self.assertIsNone(changes[2].id)

    def test_change_package_is_importable_and_formula_safe(self):
        changes = [WorkspaceChange(
            action="ADD", id=None, question="=不应成为公式", answer="+纯文本答案",
            answer_variants="", category="测试", tags="", domains="", status="approved",
            enabled="1", version=None, change_note="测试平台新增",
        )]
        payload = build_change_package(
            changes,
            {"add": 1, "update": 0, "disable": 0, "pending": 0, "ready": 1},
        )
        workbook = openpyxl.load_workbook(io.BytesIO(payload), data_only=False)
        sheet = workbook["QA内容"]
        self.assertEqual([cell.value for cell in sheet[1]], HEADERS)
        self.assertEqual(sheet["C2"].value, "=不应成为公式")
        self.assertEqual(sheet["C2"].data_type, "s")
        self.assertEqual(sheet["D2"].data_type, "s")


class StagingRefreshTests(unittest.TestCase):
    def test_snapshot_keeps_content_but_removes_formal_delivery_access(self):
        with tempfile.TemporaryDirectory() as directory:
            db_path = Path(directory) / "qabot-workspace.db"
            baseline_path = Path(directory) / "baseline.xlsx"
            connection = sqlite3.connect(db_path)
            connection.executescript(
                """
                CREATE TABLE sys_setting (k TEXT PRIMARY KEY, v TEXT);
                CREATE TABLE sys_user (id INTEGER PRIMARY KEY, enabled TEXT);
                CREATE TABLE dingtalk_group (
                    id INTEGER PRIMARY KEY, active TEXT, webhook_url TEXT,
                    webhook_secret TEXT, robot_code TEXT, note TEXT
                );
                CREATE TABLE broadcast_schedule (id INTEGER PRIMARY KEY, enabled TEXT);
                CREATE TABLE uploaded_file (
                    id INTEGER PRIMARY KEY, filename TEXT, public_url TEXT,
                    dingtalk_media_id TEXT, dingtalk_media_uploaded_at TEXT
                );
                CREATE TABLE qa_item (
                    id INTEGER PRIMARY KEY, question TEXT, answer TEXT,
                    answer_variants TEXT, category TEXT, tags TEXT, domains TEXT,
                    status TEXT, enabled TEXT, deleted TEXT, version INTEGER
                );
                INSERT INTO sys_setting VALUES ('dingtalk_client_secret', 'formal-secret');
                INSERT INTO sys_setting VALUES ('daily_brief_enabled', 'True');
                INSERT INTO sys_setting VALUES ('bot_persona', '完整内容保留');
                INSERT INTO sys_user VALUES (1, '1');
                INSERT INTO dingtalk_group VALUES (1, '1', 'https://formal', 'SEC-formal', 'robot', '正式群');
                INSERT INTO broadcast_schedule VALUES (1, '1');
                INSERT INTO uploaded_file VALUES (1, 'a.png', 'https://formal/a.png', 'media', '2026-01-01');
                INSERT INTO qa_item VALUES (1, '问题', '答案', '', '分类', '', '', 'approved', '1', '0', 3);
                """
            )
            connection.commit()
            connection.close()

            sanitize_database(db_path)
            count = write_baseline_xlsx(db_path, baseline_path)

            connection = sqlite3.connect(db_path)
            settings = dict(connection.execute("SELECT k, v FROM sys_setting"))
            group = connection.execute(
                "SELECT active, webhook_url, webhook_secret, robot_code FROM dingtalk_group"
            ).fetchone()
            schedule = connection.execute("SELECT enabled FROM broadcast_schedule").fetchone()[0]
            user = connection.execute("SELECT enabled FROM sys_user").fetchone()[0]
            upload = connection.execute(
                "SELECT public_url, dingtalk_media_id FROM uploaded_file"
            ).fetchone()
            connection.close()

            self.assertEqual(settings["dingtalk_client_secret"], "")
            self.assertEqual(settings["daily_brief_enabled"], "False")
            self.assertEqual(settings["bot_persona"], "完整内容保留")
            self.assertEqual(group, ("0", None, None, None))
            self.assertEqual(schedule, "0")
            self.assertEqual(user, "0")
            self.assertEqual(upload, ("/files/a.png", None))
            self.assertEqual(count, 1)
            workbook = openpyxl.load_workbook(baseline_path)
            self.assertEqual(workbook["QA内容"]["B2"].value, 1)


if __name__ == "__main__":
    unittest.main()
