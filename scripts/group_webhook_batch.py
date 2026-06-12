#!/usr/bin/env python3
"""
群 @全体 webhook 批量配置工具。

两个动作：
  export  → 导出所有群到 Excel，你填 webhook + secret
  import  → 把填好的 Excel 导回 DB

用法（在 backend 目录）：
  # 1. 导出（生成 ~/Desktop/群@全体配置.xlsx）
  .venv/bin/python ../scripts/group_webhook_batch.py export

  # 2. 用 Excel 打开填好 webhook + secret 列

  # 3. 导回
  .venv/bin/python ../scripts/group_webhook_batch.py import

Excel 列：
  A: 群ID（不要改）
  B: 群名（不要改，仅供识别）
  C: webhook地址（粘贴 https://oapi.dingtalk.com/robot/send?access_token=xxx）
  D: 加签密钥（粘贴 SEC 开头的密钥）
  E: 当前状态（不要改）

规则：
  - C、D 都填了 → 配置该群 @全体
  - C、D 留空 → 跳过（不动该群现有配置）
  - C 填 "CLEAR" → 清除该群 @全体 配置
"""
from __future__ import annotations
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent / "backend"
sys.path.insert(0, str(BACKEND))

from app.db import SessionLocal, DingtalkGroup  # noqa: E402
import openpyxl  # noqa: E402

XLSX = Path.home() / "Desktop" / "群@全体配置.xlsx"


def do_export():
    db = SessionLocal()
    try:
        rows = db.query(DingtalkGroup).order_by(DingtalkGroup.id).all()
        data = [(r.id, r.conversation_title or "(未命名)",
                 r.webhook_url or "", r.webhook_secret or "",
                 "✅已配" if (r.webhook_url or "").strip() else "—") for r in rows]
    finally:
        db.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "群webhook配置"
    ws.append(["群ID(勿改)", "群名(勿改)", "webhook地址", "加签密钥(SEC开头)", "当前状态(勿改)"])
    for col, w in [("A", 12), ("B", 40), ("C", 70), ("D", 60), ("E", 12)]:
        ws.column_dimensions[col].width = w
    for r in data:
        ws.append(list(r))
    wb.save(XLSX)
    print(f"✅ 已导出 {len(data)} 个群到: {XLSX}")
    print("\n填写说明：")
    print("  - C列填 webhook 地址、D列填 SEC 密钥 → 配置该群 @全体")
    print("  - C、D 留空 → 跳过该群（不改现有配置）")
    print("  - C列填 CLEAR → 清除该群 @全体 配置")
    print("\n填好后跑：.venv/bin/python ../scripts/group_webhook_batch.py import")


def do_import():
    if not XLSX.exists():
        print(f"❌ 找不到 {XLSX}，先跑 export")
        sys.exit(1)
    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active

    db = SessionLocal()
    try:
        configured = cleared = skipped = notfound = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            try:
                gid = int(row[0])
            except (TypeError, ValueError):
                continue
            url = (str(row[2]).strip() if len(row) > 2 and row[2] else "")
            secret = (str(row[3]).strip() if len(row) > 3 and row[3] else "")

            r = db.get(DingtalkGroup, gid)
            if r is None:
                notfound += 1
                continue

            if url.upper() == "CLEAR":
                r.webhook_url = None
                r.webhook_secret = None
                cleared += 1
                print(f"  🧹 清除 群{gid} {r.conversation_title}")
                continue

            if not url:
                skipped += 1
                continue

            if not url.startswith("https://oapi.dingtalk.com/robot/send"):
                print(f"  ⚠️  群{gid} webhook 格式不对，跳过: {url[:50]}")
                skipped += 1
                continue

            r.webhook_url = url[:500]
            r.webhook_secret = secret[:255] if secret else None
            configured += 1
            print(f"  ✅ 配置 群{gid} {r.conversation_title}")
        db.commit()
        print(f"\n统计：配置 {configured} / 清除 {cleared} / 跳过 {skipped} / 未找到 {notfound}")
    finally:
        db.close()


if __name__ == "__main__":
    action = sys.argv[1] if len(sys.argv) > 1 else ""
    if action == "export":
        do_export()
    elif action == "import":
        do_import()
    else:
        print("用法: group_webhook_batch.py [export|import]")
        sys.exit(1)
