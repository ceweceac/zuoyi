"""机器人单聊主动推送：给指定钉钉用户私信通知。

与群发（broadcaster.py 的自定义机器人 webhook）是两套不同机制：
- 群发：webhook + 加签，无需 accessToken。
- 私信：必须用 accessToken + robotCode，调钉钉新版 OpenAPI
  POST https://api.dingtalk.com/v1.0/robot/oToMessages/batchSend
  （token 走 header `x-acs-dingtalk-access-token`，userIds 最多 20 个/批）。

目标用户用「手机号」定位：手机号 → userId 走 topapi/v2/user/getbymobile
（需开通「根据手机号查询用户」通讯录权限）。也兼容直接填 staffId（非纯数字/非
11 位的当作 staffId 原样用）。

前置依赖（钉钉开发者后台）：
- 「机器人发送单聊消息」权限（否则 batchSend 403）。
- 用手机号定位时还需「根据手机号查询用户」通讯录权限。
"""
import io
import json
import logging
import re
import time
from typing import List, Tuple

import httpx

from ..config import settings
# 复用 media 模块的 accessToken 缓存+锁，不重复造
from .dingtalk_media import _get_access_token

log = logging.getLogger(__name__)

_BATCH_SEND_URL = "https://api.dingtalk.com/v1.0/robot/oToMessages/batchSend"
_GET_BY_MOBILE_URL = "https://oapi.dingtalk.com/topapi/v2/user/getbymobile"

# 钉钉单聊批量发一次最多 20 个 userId
_MAX_USERS_PER_CALL = 20
# 大陆手机号
_PHONE_RE = re.compile(r"1[3-9]\d{9}")


def _robot_code() -> str:
    """robotCode：优先用配置的 dingtalk_robot_code，留空回退 AppKey(client_id)。"""
    return (getattr(settings, "dingtalk_robot_code", "") or "").strip() \
        or (getattr(settings, "dingtalk_client_id", "") or "").strip()


def _excluded_staff_ids() -> set:
    """当前被隐藏（排除）的 staffId 集合。"""
    from ..db import SessionLocal, DirectPushExcluded
    db = SessionLocal()
    try:
        return {r[0] for r in db.query(DirectPushExcluded.staff_id).all()}
    finally:
        db.close()


def list_known_users(limit: int = 1000) -> List[dict]:
    """列出和机器人私聊过的用户（staffId + 姓名 + 最近活跃 + 消息数）。

    这些 staffId 由 bot.py 在私聊时存入 conversation 表，可直接用于单聊推送，
    **无需「根据手机号查询用户」通讯录权限**。
    群聊里的 sender（g:conv:user 形式）排除，被隐藏（排除名单）的也排除。
    返回按最近活跃倒序。
    """
    from sqlalchemy import func
    from ..db import SessionLocal, Conversation
    db = SessionLocal()
    try:
        rows = (
            db.query(
                Conversation.sender,
                func.max(Conversation.sender_name),
                func.count(Conversation.id),
                func.max(Conversation.created_at),
            )
            .filter(Conversation.sender.isnot(None), Conversation.sender != "")
            .group_by(Conversation.sender)
            .all()
        )
    finally:
        db.close()
    excluded = _excluded_staff_ids()
    out = []
    for sender, name, cnt, last_at in rows:
        if not sender or sender.startswith("g:"):
            continue  # 群聊上下文的伪 sender，跳过
        # 只保留像真实钉钉 staffId 的（字母/数字/下划线/连字符），
        # 过滤早期直接调 pipeline 测试时写入的中文问题文本等脏 sender
        if not re.fullmatch(r"[A-Za-z0-9_\-]+", sender):
            continue
        if sender in excluded:
            continue  # 已隐藏（离职/测试/无效）
        out.append({
            "staff_id": sender,
            "name": (name or "").strip() or "(未知姓名)",
            "msg_count": int(cnt or 0),
            "last_at": last_at.strftime("%Y-%m-%d %H:%M") if last_at else "",
            "_sort": last_at or "",
        })
    out.sort(key=lambda x: x["_sort"], reverse=True)
    for x in out:
        x.pop("_sort", None)
    return out[:limit]


def hide_users(staff_ids: List[str], reason: str = "", by: str = "",
               names: dict = None) -> int:
    """把一批 staffId 加入排除名单（隐藏，不删对话记录）。返回新增条数。"""
    from ..db import SessionLocal, DirectPushExcluded
    names = names or {}
    ids = [s.strip() for s in (staff_ids or []) if s and s.strip()]
    if not ids:
        return 0
    db = SessionLocal()
    added = 0
    try:
        existing = {r[0] for r in db.query(DirectPushExcluded.staff_id)
                    .filter(DirectPushExcluded.staff_id.in_(ids)).all()}
        for sid in ids:
            if sid in existing:
                continue
            db.add(DirectPushExcluded(
                staff_id=sid, name=(names.get(sid) or "")[:128],
                reason=(reason or "")[:255], created_by=by or "",
            ))
            added += 1
        db.commit()
    except Exception:
        log.exception("hide_users 失败")
        db.rollback()
    finally:
        db.close()
    return added


def unhide_users(staff_ids: List[str]) -> int:
    """从排除名单移除（恢复）。返回移除条数。"""
    from ..db import SessionLocal, DirectPushExcluded
    ids = [s.strip() for s in (staff_ids or []) if s and s.strip()]
    if not ids:
        return 0
    db = SessionLocal()
    try:
        n = db.query(DirectPushExcluded).filter(
            DirectPushExcluded.staff_id.in_(ids)).delete(synchronize_session=False)
        db.commit()
        return n
    except Exception:
        log.exception("unhide_users 失败")
        db.rollback()
        return 0
    finally:
        db.close()


def list_excluded() -> List[dict]:
    """列出所有被隐藏的用户。"""
    from ..db import SessionLocal, DirectPushExcluded
    db = SessionLocal()
    try:
        rows = db.query(DirectPushExcluded).order_by(
            DirectPushExcluded.created_at.desc()).all()
        return [{
            "staff_id": r.staff_id,
            "name": r.name or "(未知)",
            "reason": r.reason or "",
            "by": r.created_by or "",
            "at": r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else "",
        } for r in rows]
    finally:
        db.close()


# ─────────── 可复用联系人名单（各位老师等）───────────

def list_contacts() -> List[dict]:
    """列出所有已存联系人。按姓名排序。"""
    from ..db import SessionLocal, DirectContact
    db = SessionLocal()
    try:
        rows = db.query(DirectContact).order_by(DirectContact.name).all()
        return [{
            "id": r.id, "staff_id": r.staff_id, "name": r.name,
            "tag": r.tag or "",
            "at": r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else "",
        } for r in rows]
    finally:
        db.close()


def add_contacts(items: List[dict], by: str = "") -> int:
    """新增联系人。items: [{"staff_id":..,"name":..,"tag":..}]。
    staff_id 已存在则更新姓名/标签。返回新增+更新条数。
    """
    from ..db import SessionLocal, DirectContact
    n = 0
    db = SessionLocal()
    try:
        for it in items or []:
            sid = (it.get("staff_id") or "").strip()
            name = (it.get("name") or "").strip()
            if not sid or not name:
                continue
            row = db.query(DirectContact).filter(
                DirectContact.staff_id == sid).first()
            if row:
                row.name = name[:128]
                row.tag = (it.get("tag") or "")[:64] or row.tag
            else:
                db.add(DirectContact(
                    staff_id=sid[:64], name=name[:128],
                    tag=(it.get("tag") or "")[:64], created_by=by or "",
                ))
            n += 1
        db.commit()
    except Exception:
        log.exception("add_contacts 失败")
        db.rollback()
    finally:
        db.close()
    return n


def remove_contacts(staff_ids: List[str]) -> int:
    """删除联系人（按 staffId）。返回删除条数。"""
    from ..db import SessionLocal, DirectContact
    ids = [s.strip() for s in (staff_ids or []) if s and s.strip()]
    if not ids:
        return 0
    db = SessionLocal()
    try:
        n = db.query(DirectContact).filter(
            DirectContact.staff_id.in_(ids)).delete(synchronize_session=False)
        db.commit()
        return n
    except Exception:
        log.exception("remove_contacts 失败")
        db.rollback()
        return 0
    finally:
        db.close()


def extract_phones(raw_text: str) -> List[str]:
    """从一段文本里抽出所有大陆手机号并去重（保持出现顺序）。

    支持带空格/连字符的写法（如「139 9999 8888」「139-9999-8888」）——
    按行/分隔切分后，逐段去掉内部空格和连字符再匹配。
    """
    if not raw_text:
        return []
    seen = set()
    out = []
    # 按行/常见分隔切，逐段规整空格和短横线
    for seg in re.split(r"[\n\r\t,，;；]+", raw_text):
        cleaned = re.sub(r"[\s\-]+", "", seg)
        for m in _PHONE_RE.findall(cleaned):
            if m not in seen:
                seen.add(m)
                out.append(m)
    return out


def parse_sheet(file_bytes: bytes, filename: str) -> List[str]:
    """从上传的表格里抽手机号。支持 .xlsx（openpyxl）/ .csv / .txt（按文本抽号）。
    遍历所有单元格，正则抽 1[3-9]\\d{9}，去重。
    """
    name = (filename or "").lower()
    texts: List[str] = []
    try:
        if name.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    for cell in row:
                        if cell is not None:
                            texts.append(str(cell))
            wb.close()
        else:
            # csv / txt / 其它：当纯文本解码后整体抽号
            try:
                texts.append(file_bytes.decode("utf-8-sig"))
            except UnicodeDecodeError:
                texts.append(file_bytes.decode("gbk", errors="ignore"))
    except Exception as e:
        log.exception("parse_sheet failed: %s", filename)
        raise RuntimeError(f"解析表格失败：{e}")
    return extract_phones("\n".join(texts))


def _get_userid_by_mobile(mobile: str, token: str) -> Tuple[str, str]:
    """手机号 → (userId, name)。查不到/无权限返回 ("", "")。"""
    try:
        r = httpx.post(
            f"{_GET_BY_MOBILE_URL}?access_token={token}",
            json={"mobile": mobile}, timeout=10,
        )
        data = r.json()
    except Exception as e:
        log.warning("getbymobile %s 异常: %s", mobile, e)
        return "", ""
    if data.get("errcode") != 0:
        log.warning("getbymobile %s 失败: %s", mobile, data)
        return "", ""
    result = data.get("result") or {}
    return result.get("userid") or "", result.get("name") or ""


def _load_mobile_cache(mobiles: List[str]) -> dict:
    """批量读手机号缓存，返回 {mobile: userid}。"""
    if not mobiles:
        return {}
    from ..db import SessionLocal, DingtalkUserCache
    db = SessionLocal()
    try:
        rows = db.query(DingtalkUserCache.mobile, DingtalkUserCache.userid)\
            .filter(DingtalkUserCache.mobile.in_(mobiles)).all()
        return {m: u for m, u in rows if u}
    finally:
        db.close()


def _save_mobile_cache(mobile: str, userid: str, name: str = "") -> None:
    """写/更新手机号缓存。失败静默。"""
    from ..db import SessionLocal, DingtalkUserCache
    db = SessionLocal()
    try:
        row = db.query(DingtalkUserCache).filter(
            DingtalkUserCache.mobile == mobile).first()
        if row:
            row.userid = userid
            if name:
                row.name = name[:128]
        else:
            db.add(DingtalkUserCache(mobile=mobile, userid=userid, name=name[:128]))
        db.commit()
    except Exception:
        log.exception("写手机号缓存失败: %s", mobile)
        db.rollback()
    finally:
        db.close()


def resolve_userids(targets: List[str]) -> Tuple[List[str], List[str], dict]:
    """把目标列表（手机号或 staffId）解析成 userId 列表。

    手机号先查本地缓存（dingtalk_user_cache），命中则零 HTTP；未命中才调
    getbymobile，成功后写回缓存。大批量重复推送时几乎不再调钉钉接口。

    返回 (userids, unresolved, mobile_map)：
    - userids：成功拿到的 userId（去重）
    - unresolved：查不到 userId 的手机号（用于回报）
    - mobile_map：手机号 → userId，便于结果展示
    """
    userids: List[str] = []
    unresolved: List[str] = []
    mobile_map: dict = {}
    seen = set()

    # 先分出手机号，批量查缓存
    mobiles = [t.strip() for t in targets if t and _PHONE_RE.fullmatch(t.strip())]
    cache = _load_mobile_cache(mobiles)
    token = None  # 懒加载：只有需要调接口时才取

    for t in targets:
        t = (t or "").strip()
        if not t:
            continue
        if _PHONE_RE.fullmatch(t):
            uid = cache.get(t)
            if not uid:
                if token is None:
                    token = _get_access_token()
                uid, name = _get_userid_by_mobile(t, token)
                if not uid:
                    unresolved.append(t)
                    continue
                _save_mobile_cache(t, uid, name)
            mobile_map[t] = uid
        else:
            uid = t  # 当 staffId 原样用
        if uid not in seen:
            seen.add(uid)
            userids.append(uid)
    return userids, unresolved, mobile_map


def _build_msg(content: str, msg_type: str) -> Tuple[str, str]:
    """构造 (msgKey, msgParam)。msg_type ∈ {text, markdown}。"""
    if msg_type == "markdown":
        param = {"title": "通知", "text": content}
        return "sampleMarkdown", json.dumps(param, ensure_ascii=False)
    return "sampleText", json.dumps({"content": content}, ensure_ascii=False)


def _batch_send(userids: List[str], msg_key: str, msg_param: str, token: str) -> dict:
    """调 batchSend 发一批（≤20 人）。返回 {ok, invalid:[...], error?, response}。"""
    payload = {
        "robotCode": _robot_code(),
        "userIds": userids,
        "msgKey": msg_key,
        "msgParam": msg_param,
    }
    try:
        r = httpx.post(
            _BATCH_SEND_URL,
            headers={"x-acs-dingtalk-access-token": token,
                     "Content-Type": "application/json"},
            json=payload, timeout=20,
        )
        try:
            data = r.json()
        except Exception:
            data = {"status_code": r.status_code, "body": r.text[:300]}
    except Exception as e:
        log.exception("batchSend 异常")
        return {"ok": False, "error": str(e), "invalid": []}

    # 成功响应含 processQueryKey；失败 HTTP 4xx 带 code/message
    if r.status_code == 200 and isinstance(data, dict) and "processQueryKey" in data:
        # invalidStaffIdList：钉钉判定无效的 userId（如不在可见范围）
        invalid = data.get("invalidStaffIdList") or data.get("flowControlledStaffIdList") or []
        return {"ok": True, "invalid": invalid, "response": data}
    log.warning("batchSend 失败: status=%s body=%s", r.status_code, data)
    msg = data.get("message") if isinstance(data, dict) else str(data)
    return {"ok": False, "error": f"HTTP {r.status_code}: {msg}", "invalid": [],
            "response": data}


def start_push_async(targets: List[str], content: str, msg_type: str = "text",
                     created_by: str = "") -> dict:
    """建一条 pending 的推送任务并丢到后台执行，立即返回 push_id。

    UI 调它后立即返回，发送进度/结果在 direct_push 表里轮询。
    返回 {ok, push_id} 或 {ok:False, error}。
    """
    if not _robot_code():
        return {"ok": False, "error": "未配置 robotCode 且 dingtalk_client_id 为空，无法推送"}
    content = (content or "").strip()
    if not content:
        return {"ok": False, "error": "推送内容为空"}
    targets = [t.strip() for t in (targets or []) if t and t.strip()]
    # 去重，保持顺序
    targets = list(dict.fromkeys(targets))
    if not targets:
        return {"ok": False, "error": "未指定任何目标用户"}

    from ..db import SessionLocal, DirectPush
    db = SessionLocal()
    try:
        row = DirectPush(
            content_text=content[:5000],
            msg_type=msg_type,
            target_count=len(targets),
            user_count=0,
            sent_count=0,
            status="pending",
            result_summary="排队中…",
            raw_targets=json.dumps(targets, ensure_ascii=False)[:100000],
            created_by=created_by or "",
        )
        db.add(row)
        db.commit()
        push_id = row.id
    except Exception as e:
        log.exception("建私信推送任务失败")
        db.rollback()
        return {"ok": False, "error": f"创建任务失败：{e}"}
    finally:
        db.close()

    # 丢后台执行；scheduler 没起来则当场同步跑（兜底）
    from . import scheduler
    queued = scheduler.run_in_background(run_push, push_id,
                                         job_id=f"direct_push_{push_id}")
    if not queued:
        log.warning("scheduler 未就绪，私信任务 #%s 同步执行", push_id)
        run_push(push_id)
    return {"ok": True, "push_id": push_id}


def run_push(push_id: int) -> None:
    """后台执行一条私信推送任务：解析 userId → 分批发送 → 实时更新进度。"""
    from ..db import SessionLocal, DirectPush

    db = SessionLocal()
    try:
        row = db.get(DirectPush, push_id)
        if row is None:
            log.warning("run_push: #%s 不存在", push_id)
            return
        try:
            targets = json.loads(row.raw_targets or "[]")
        except Exception:
            targets = []
        content = row.content_text or ""
        msg_type = row.msg_type or "text"
        row.status = "running"
        row.result_summary = "解析用户中…"
        db.commit()
    finally:
        db.close()

    if not targets:
        _finish(push_id, 0, 0, [], [], ["任务无目标"])
        return

    try:
        token = _get_access_token()
    except Exception as e:
        _finish(push_id, 0, 0, [], [], [f"获取 access_token 失败：{e}"])
        return

    try:
        userids, unresolved, _ = resolve_userids(targets)
    except Exception as e:
        _finish(push_id, 0, 0, [], [], [f"解析用户失败：{e}"])
        return

    if not userids:
        _finish(push_id, 0, 0, unresolved, [], ["没有可推送的有效用户"])
        return

    # 落 user_count，让 UI 进度分母可见
    _update_progress(push_id, user_count=len(userids), sent_count=0,
                     summary=f"发送中… 0/{len(userids)}")

    sent, invalid, failed_userids, errors = _send_in_batches(
        userids, content, msg_type, token, push_id)

    _finish(push_id, sent, len(userids), unresolved, invalid, errors,
            failed_userids=failed_userids)


def _send_in_batches(userids, content, msg_type, token, push_id=None):
    """分批发送，返回 (sent, invalid, failed_userids, errors)。
    每批结束更新进度（若给了 push_id）。
    """
    msg_key, msg_param = _build_msg(content, msg_type)
    sent = 0
    invalid: List[str] = []
    failed_userids: List[str] = []
    errors: List[str] = []
    total = len(userids)
    for i in range(0, total, _MAX_USERS_PER_CALL):
        chunk = userids[i:i + _MAX_USERS_PER_CALL]
        res = _batch_send(chunk, msg_key, msg_param, token)
        if res.get("ok"):
            bad = set(res.get("invalid") or [])
            invalid.extend(bad)
            sent += len(chunk) - len(bad)
        else:
            errors.append(res.get("error") or "发送失败")
            failed_userids.extend(chunk)  # 整批失败，全部计入待重试
        if push_id is not None:
            _update_progress(push_id, sent_count=sent,
                             summary=f"发送中… {min(i + len(chunk), total)}/{total}")
        time.sleep(0.3)
    return sent, invalid, failed_userids, errors


def retry_push(push_id: int) -> dict:
    """重试一条推送里失败的部分（failed_userids + unresolved 手机号）。
    返回 {ok, push_id}（重试也异步跑）或 {ok:False, error}。
    """
    from ..db import SessionLocal, DirectPush
    db = SessionLocal()
    try:
        row = db.get(DirectPush, push_id)
        if row is None:
            return {"ok": False, "error": "推送记录不存在"}
        try:
            failed = json.loads(row.failed_userids or "[]")
        except Exception:
            failed = []
        try:
            unresolved = json.loads(row.unresolved or "[]")
        except Exception:
            unresolved = []
        content = row.content_text or ""
        msg_type = row.msg_type or "text"
        by = row.created_by or ""
    finally:
        db.close()

    retry_targets = list(dict.fromkeys([*failed, *unresolved]))
    if not retry_targets:
        return {"ok": False, "error": "没有需要重试的失败项"}
    # 新建一条重试任务（独立记录，便于追溯）
    return start_push_async(retry_targets, content, msg_type, created_by=by)


def _update_progress(push_id: int, user_count: int = None,
                     sent_count: int = None, summary: str = None) -> None:
    """更新进度字段（running 中频繁调，失败静默）。"""
    from ..db import SessionLocal, DirectPush
    db = SessionLocal()
    try:
        row = db.get(DirectPush, push_id)
        if row is None:
            return
        if user_count is not None:
            row.user_count = user_count
        if sent_count is not None:
            row.sent_count = sent_count
        if summary is not None:
            row.result_summary = summary[:1000]
        db.commit()
    except Exception:
        db.rollback()
    finally:
        db.close()


def _finish(push_id: int, sent: int, user_count: int, unresolved: list,
            invalid: list, errors: list, failed_userids: list = None) -> None:
    """收尾：写最终状态 + 结果摘要。"""
    from ..db import SessionLocal, DirectPush
    failed_userids = failed_userids or []
    if sent == 0 and (errors or not user_count):
        status = "failed"
    elif errors or unresolved or invalid:
        status = "partial"
    else:
        status = "sent"
    parts = [f"成功 {sent} 人"]
    if unresolved:
        parts.append(f"未识别手机号 {len(unresolved)}")
    if invalid:
        parts.append(f"无效/受限 {len(invalid)}")
    if errors:
        parts.append("错误：" + "；".join(errors[:3]))
    db = SessionLocal()
    try:
        row = db.get(DirectPush, push_id)
        if row is None:
            return
        row.sent_count = sent
        row.user_count = user_count
        row.unresolved = json.dumps(unresolved, ensure_ascii=False)[:4000]
        row.invalid = json.dumps(list(invalid), ensure_ascii=False)[:4000]
        row.failed_userids = json.dumps(failed_userids, ensure_ascii=False)[:8000]
        row.status = status
        row.result_summary = "；".join(parts)[:1000]
        db.commit()
    except Exception:
        log.exception("私信任务收尾写库失败 #%s", push_id)
        db.rollback()
    finally:
        db.close()
