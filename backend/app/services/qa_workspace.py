"""隔离 QA 测试工作台：基线导入、审核差异统计和正式环境变更包导出。"""
from __future__ import annotations

import io
import json
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from ..config import settings
from ..db import QaItem, SessionLocal, SysUser, hash_password


HEADERS = [
    "操作", "ID", "问题", "标准答案", "备用答案(JSON)", "分类",
    "标签", "业务域", "状态", "启用", "版本", "修改说明",
]
CONTENT_FIELDS = ("question", "answer", "answer_variants", "category", "tags", "domains")


@dataclass(frozen=True)
class BaselineQa:
    id: int
    question: str
    answer: str
    answer_variants: str
    category: str
    tags: str
    domains: str
    status: str
    enabled: str
    version: int


@dataclass(frozen=True)
class WorkspaceChange:
    action: str
    id: int | None
    question: str
    answer: str
    answer_variants: str
    category: str
    tags: str
    domains: str
    status: str
    enabled: str
    version: int | None
    change_note: str

    def as_row(self) -> list[object]:
        return [
            self.action, self.id, self.question, self.answer, self.answer_variants,
            self.category, self.tags, self.domains, self.status, self.enabled,
            self.version, self.change_note,
        ]


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _int(value, default: int = 0) -> int:
    try:
        return int(float(str(value)))
    except (TypeError, ValueError):
        return default


def normalize_variants(value) -> str:
    """把 JSON 数组或逐行备用答案规范为去重后的 JSON。"""
    text = _text(value)
    if not text:
        return ""
    try:
        values = json.loads(text) if text.startswith("[") else text.splitlines()
    except json.JSONDecodeError as exc:
        raise ValueError("备用答案 JSON 格式不正确") from exc
    if not isinstance(values, list) or not all(isinstance(item, str) for item in values):
        raise ValueError("备用答案必须是字符串数组，或每行一个答案")
    unique: list[str] = []
    for item in values:
        item = item.strip()
        if item and item not in unique:
            unique.append(item)
    return json.dumps(unique, ensure_ascii=False) if unique else ""


def variants_as_lines(value) -> str:
    normalized = normalize_variants(value)
    return "\n".join(json.loads(normalized)) if normalized else ""


def baseline_path() -> Path:
    configured = settings.qa_workspace_baseline_xlsx.strip()
    if configured:
        return Path(configured).expanduser().resolve()
    # qa_workspace.py -> services -> app -> backend -> repository
    return Path(__file__).resolve().parents[3] / "qa" / "qa_content.xlsx"


def load_baseline(path: Path | None = None) -> dict[int, BaselineQa]:
    source = path or baseline_path()
    if not source.exists():
        raise RuntimeError(f"QA 测试基线不存在：{source}")
    workbook = openpyxl.load_workbook(source, read_only=False, data_only=False)
    if "QA内容" not in workbook.sheetnames:
        raise RuntimeError("QA 测试基线缺少“QA内容”工作表")
    sheet = workbook["QA内容"]
    header_row = None
    for cells in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 12, 12), max_col=len(HEADERS)):
        if [_text(cell.value) for cell in cells] == HEADERS:
            header_row = cells[0].row
            break
    if header_row is None:
        raise RuntimeError("QA 测试基线表头不匹配")

    rows: dict[int, BaselineQa] = {}
    for cells in sheet.iter_rows(min_row=header_row + 1, max_col=len(HEADERS)):
        values = [cell.value for cell in cells]
        qa_id = _int(values[1])
        question = _text(values[2])
        answer = _text(values[3])
        if not qa_id and not question and not answer:
            continue
        if not qa_id:
            # 基线只追踪已有正式 QA；无 ID 的 ADD 预留行不是基线。
            continue
        if qa_id in rows:
            raise RuntimeError(f"QA 测试基线 ID 重复：{qa_id}")
        rows[qa_id] = BaselineQa(
            id=qa_id,
            question=question,
            answer=answer,
            answer_variants=normalize_variants(values[4]),
            category=_text(values[5]),
            tags=_text(values[6]),
            domains=_text(values[7]),
            status=_text(values[8]) or "approved",
            enabled=_text(values[9]) or "1",
            version=_int(values[10], 1) or 1,
        )
    if not rows:
        raise RuntimeError("QA 测试基线没有有效数据")
    return rows


def ensure_workspace_ready() -> None:
    """首次启动时导入基线，并只启用 owner/partner 两个强密码账号。"""
    owner = settings.qa_workspace_owner_username.strip()
    partner = settings.qa_workspace_partner_username.strip()
    owner_password = settings.qa_workspace_owner_password
    partner_password = settings.qa_workspace_partner_password
    partner_role = settings.qa_workspace_partner_role.strip().lower()
    if not owner or not partner or owner == partner:
        raise RuntimeError("QA 测试平台 owner/partner 用户名不能为空且不能相同")
    if len(owner_password) < 12 or len(partner_password) < 12:
        raise RuntimeError("QA 测试平台账号密码必须至少 12 位，禁止使用默认弱密码")
    if partner_role not in {"editor", "admin"}:
        raise RuntimeError("QA 测试平台伙伴角色只能是 editor 或 admin")

    baseline = load_baseline()
    db = SessionLocal()
    try:
        if db.query(QaItem).count() == 0:
            now = datetime.utcnow()
            db.add_all([
                QaItem(
                    id=row.id,
                    question=row.question,
                    answer=row.answer,
                    answer_variants=row.answer_variants or None,
                    category=row.category or None,
                    tags=row.tags or None,
                    domains=row.domains or None,
                    status="approved",
                    enabled="1",
                    deleted="0",
                    version=row.version,
                    created_by="qa-workspace-baseline",
                    created_at=now,
                    updated_by="qa-workspace-baseline",
                    updated_at=now,
                    approved_by="qa-workspace-baseline",
                    approved_at=now,
                )
                for row in baseline.values()
            ])

        allowed = {owner, partner}
        db.query(SysUser).filter(~SysUser.username.in_(allowed)).update(
            {"enabled": "0"}, synchronize_session=False,
        )
        account_specs = [
            (owner, owner_password, "QA 审核负责人", "admin"),
            (partner, partner_password, "测试环境开发伙伴", partner_role),
        ]
        for username, password, display_name, role in account_specs:
            account = db.query(SysUser).filter(SysUser.username == username).first()
            if account is None:
                account = SysUser(username=username)
                db.add(account)
            account.password = hash_password(password)
            account.display_name = display_name
            account.role = role
            account.enabled = "1"
        db.commit()
    finally:
        db.close()


def _current_value(item: QaItem, name: str) -> str:
    return _text(getattr(item, name, None))


def collect_changes(
    current_items: Iterable[QaItem],
    baseline: dict[int, BaselineQa] | None = None,
) -> tuple[list[WorkspaceChange], dict[str, int]]:
    """对比当前测试库与正式基线；待审核内容不会进入导出包。"""
    base = baseline or load_baseline()
    current = {int(item.id): item for item in current_items}
    changes: list[WorkspaceChange] = []
    stats = {"add": 0, "update": 0, "disable": 0, "pending": 0, "ready": 0}

    for qa_id, before in sorted(base.items()):
        item = current.get(qa_id)
        if item is None or _current_value(item, "deleted") == "1" or _current_value(item, "enabled") == "0":
            source = item
            changes.append(WorkspaceChange(
                action="DISABLE",
                id=qa_id,
                question=_current_value(source, "question") if source else before.question,
                answer=_current_value(source, "answer") if source else before.answer,
                answer_variants=_current_value(source, "answer_variants") if source else before.answer_variants,
                category=_current_value(source, "category") if source else before.category,
                tags=_current_value(source, "tags") if source else before.tags,
                domains=_current_value(source, "domains") if source else before.domains,
                status="disabled",
                enabled="0",
                version=before.version,
                change_note=f"测试平台停用 · {getattr(source, 'updated_by', '') or 'qa_owner'}",
            ))
            stats["disable"] += 1
            continue

        changed = any(_current_value(item, field) != _text(getattr(before, field)) for field in CONTENT_FIELDS)
        if not changed:
            continue
        if _current_value(item, "status") != "approved":
            stats["pending"] += 1
            continue
        changes.append(WorkspaceChange(
            action="KEEP",
            id=qa_id,
            question=_current_value(item, "question"),
            answer=_current_value(item, "answer"),
            answer_variants=normalize_variants(_current_value(item, "answer_variants")),
            category=_current_value(item, "category"),
            tags=_current_value(item, "tags"),
            domains=_current_value(item, "domains"),
            status="approved",
            enabled="1",
            version=before.version,
            change_note=f"测试平台更新 · {item.updated_by or 'qa_partner'}",
        ))
        stats["update"] += 1

    for qa_id, item in sorted(current.items()):
        if qa_id in base or _current_value(item, "deleted") == "1" or _current_value(item, "enabled") == "0":
            continue
        if _current_value(item, "status") != "approved":
            stats["pending"] += 1
            continue
        changes.append(WorkspaceChange(
            action="ADD",
            id=None,
            question=_current_value(item, "question"),
            answer=_current_value(item, "answer"),
            answer_variants=normalize_variants(_current_value(item, "answer_variants")),
            category=_current_value(item, "category"),
            tags=_current_value(item, "tags"),
            domains=_current_value(item, "domains"),
            status="approved",
            enabled="1",
            version=None,
            change_note=f"测试平台新增 · {item.updated_by or item.created_by or 'qa_partner'}",
        ))
        stats["add"] += 1

    stats["ready"] = stats["add"] + stats["update"] + stats["disable"]
    return changes, stats


def workspace_changes() -> tuple[list[WorkspaceChange], dict[str, int]]:
    db = SessionLocal()
    try:
        return collect_changes(db.query(QaItem).all())
    finally:
        db.close()


def _write_text_cell(cell, value) -> None:
    """强制按文本写入，避免问题/答案以 = + - @ 开头时变成 Excel 公式。"""
    cell.value = "" if value is None else value
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        cell.data_type = "s"


def build_change_package(changes: list[WorkspaceChange], stats: dict[str, int]) -> bytes:
    if not changes:
        raise ValueError("当前没有已审核、可合并的 QA 修改")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "QA内容"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L{len(changes) + 1}"
    header_fill = PatternFill("solid", fgColor="4338CA")
    editable_fill = PatternFill("solid", fgColor="FEF3C7")
    for column, header in enumerate(HEADERS, 1):
        cell = sheet.cell(1, column, header)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_no, change in enumerate(changes, 2):
        for column, value in enumerate(change.as_row(), 1):
            cell = sheet.cell(row_no, column)
            _write_text_cell(cell, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if column in (1, 3, 4, 5, 6, 7, 8, 12):
                cell.fill = editable_fill
    widths = [11, 10, 42, 72, 34, 15, 24, 20, 12, 9, 9, 32]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 28

    guide = workbook.create_sheet("合并说明")
    guide.append(["QA 测试平台正式环境变更包"])
    guide["A1"].font = Font(size=16, bold=True, color="312E81")
    guide.append([f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    guide.append([f"可合并：新增 {stats['add']}，更新 {stats['update']}，停用 {stats['disable']}"])
    guide.append([f"未包含的待审核修改：{stats['pending']}"])
    guide.append(["请先运行 import-db 的 dry-run；确认后再加 --apply。正式库写入前会自动备份并检查版本冲突。"])
    guide.column_dimensions["A"].width = 110
    for row in guide.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_change_package() -> tuple[bytes, dict[str, int]]:
    changes, stats = workspace_changes()
    return build_change_package(changes, stats), stats
